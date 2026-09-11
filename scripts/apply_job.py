from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook


TRACKER_FILE = Path("data/job_tracker.xlsx")


HEADERS = [
    "Date Found",
    "Job Title",
    "Company",
    "Location",
    "Score",
    "Category",
    "Recommendation",
    "Role Match",
    "Matching Skills",
    "Missing Skills",
    "Warnings",
    "Posted",
    "Deadline",
    "URL",
    "Application Status",
    "CV Version",
    "Cover Letter",
    "Notes",
]


def get_column_indexes(sheet):
    """Create a mapping between column names and column numbers."""

    indexes = {}

    for cell in sheet[1]:
        indexes[cell.value] = cell.column

    return indexes


def get_jobs_to_apply():
    """Return jobs that are recommended for application/review."""

    if not TRACKER_FILE.exists():
        print("Tracker file does not exist.")
        return []

    workbook = load_workbook(TRACKER_FILE)
    sheet = workbook["Jobs"]

    columns = get_column_indexes(sheet)

    jobs = []

    for row_number in range(2, sheet.max_row + 1):

        recommendation = sheet.cell(
            row=row_number,
            column=columns["Recommendation"]
        ).value

        status = sheet.cell(
            row=row_number,
            column=columns["Application Status"]
        ).value

        # Only show jobs that have not already been applied for.
        if (
            recommendation in ["APPLY", "REVIEW"]
            and status in ["Not Applied", "To Apply"]
        ):
            jobs.append({
                "row": row_number,
                "title": sheet.cell(
                    row=row_number,
                    column=columns["Job Title"]
                ).value,
                "company": sheet.cell(
                    row=row_number,
                    column=columns["Company"]
                ).value,
                "location": sheet.cell(
                    row=row_number,
                    column=columns["Location"]
                ).value,
                "score": sheet.cell(
                    row=row_number,
                    column=columns["Score"]
                ).value,
                "category": sheet.cell(
                    row=row_number,
                    column=columns["Category"]
                ).value,
                "recommendation": recommendation,
                "role_match": sheet.cell(
                    row=row_number,
                    column=columns["Role Match"]
                ).value,
                "matching_skills": sheet.cell(
                    row=row_number,
                    column=columns["Matching Skills"]
                ).value,
                "missing_skills": sheet.cell(
                    row=row_number,
                    column=columns["Missing Skills"]
                ).value,
                "warnings": sheet.cell(
                    row=row_number,
                    column=columns["Warnings"]
                ).value,
                "deadline": sheet.cell(
                    row=row_number,
                    column=columns["Deadline"]
                ).value,
                "url": sheet.cell(
                    row=row_number,
                    column=columns["URL"]
                ).value,
            })

    workbook.close()

    return jobs


def display_jobs(jobs):
    """Display available jobs."""

    print("=" * 60)
    print("APPLICATION MANAGER")
    print("=" * 60)
    print()

    print(f"Found {len(jobs)} jobs to consider.")
    print()

    for index, job in enumerate(jobs, start=1):

        print(
            f"{index}. {job['score']}% — "
            f"{job['title']} at {job['company']}"
        )

        print(
            f"   Location: {job['location']}"
        )

        print(
            f"   Recommendation: {job['recommendation']}"
        )

        print()


def display_job_details(job):
    """Display detailed information about a job."""

    print("-" * 60)
    print(f"{job['score']}% — {job['title']}")
    print("-" * 60)

    print(f"Company: {job['company']}")
    print(f"Location: {job['location']}")
    print(f"Category: {job['category']}")
    print(f"Recommendation: {job['recommendation']}")
    print(f"Role match: {job['role_match'] or 'Not specified'}")

    if job["matching_skills"]:
        print(
            f"Matching skills: "
            f"{job['matching_skills']}"
        )

    if job["missing_skills"]:
        print(
            f"Missing skills: "
            f"{job['missing_skills']}"
        )

    if job["warnings"]:
        print(
            f"Warnings: "
            f"{job['warnings']}"
        )

    print(
        f"Deadline: "
        f"{job['deadline'] or 'Not specified'}"
    )

    print()
    print(f"Application URL:")
    print(job["url"])

    print("-" * 60)


def update_application(job):
    """Update the selected job with application information."""

    workbook = load_workbook(TRACKER_FILE)
    sheet = workbook["Jobs"]

    columns = get_column_indexes(sheet)

    row = job["row"]

    print()
    print("APPLICATION DETAILS")
    print()

    cv_version = input(
        "CV version used "
        "(e.g. Backend CV v1): "
    ).strip()

    while not cv_version:
        print("Please enter a CV version.")
        cv_version = input(
            "CV version used: "
        ).strip()

    cover_letter = input(
        "Cover letter used? [y/n]: "
    ).strip().lower()

    while cover_letter not in ["y", "n"]:
        cover_letter = input(
            "Please enter y or n: "
        ).strip().lower()

    cover_letter_value = (
        "Yes"
        if cover_letter == "y"
        else "No"
    )

    notes = input(
        "Notes "
        "(press Enter to skip): "
    ).strip()

    # Update tracker
    sheet.cell(
        row=row,
        column=columns["Application Status"]
    ).value = "Applied"

    sheet.cell(
        row=row,
        column=columns["CV Version"]
    ).value = cv_version

    sheet.cell(
        row=row,
        column=columns["Cover Letter"]
    ).value = cover_letter_value

    sheet.cell(
        row=row,
        column=columns["Notes"]
    ).value = notes

    # Add application date to Notes for now.
    application_date = datetime.now().strftime(
        "%Y-%m-%d"
    )

    existing_notes = sheet.cell(
        row=row,
        column=columns["Notes"]
    ).value

    date_note = (
        f"Applied: {application_date}"
    )

    if existing_notes:
        sheet.cell(
            row=row,
            column=columns["Notes"]
        ).value = (
            f"{date_note} | {existing_notes}"
        )
    else:
        sheet.cell(
            row=row,
            column=columns["Notes"]
        ).value = date_note

    workbook.save(TRACKER_FILE)
    workbook.close()

    print()
    print("=" * 60)
    print("APPLICATION RECORDED")
    print("=" * 60)
    print()
    print(f"Job: {job['title']}")
    print(f"Company: {job['company']}")
    print(f"Status: Applied")
    print(f"CV: {cv_version}")
    print(f"Cover Letter: {cover_letter_value}")
    print(f"Date: {application_date}")
    print()


def main():

    jobs = get_jobs_to_apply()

    if not jobs:
        print("=" * 60)
        print("APPLICATION MANAGER")
        print("=" * 60)
        print()
        print("No jobs currently require application.")
        return

    display_jobs(jobs)

    print(
        "Enter the number of the job you want to process."
    )
    print("Enter 0 to exit.")
    print()

    while True:

        choice = input(
            "Select job: "
        ).strip()

        try:
            choice = int(choice)
        except ValueError:
            print("Please enter a number.")
            continue

        if choice == 0:
            print("Exiting.")
            return

        if 1 <= choice <= len(jobs):
            break

        print(
            f"Please enter a number between "
            f"0 and {len(jobs)}."
        )

    selected_job = jobs[choice - 1]

    print()

    display_job_details(selected_job)

    print()
    confirm = input(
        "Have you reviewed this job and want to record an application? "
        "[y/n]: "
    ).strip().lower()

    if confirm != "y":
        print()
        print("Application not recorded.")
        return

    update_application(selected_job)


if __name__ == "__main__":
    main()