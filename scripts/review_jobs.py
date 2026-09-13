
from pathlib import Path

from openpyxl import load_workbook


TRACKER_FILE = Path("data/job_tracker.xlsx")


# ==========================================================
# LOAD JOBS
# ==========================================================

def load_jobs():
    """
    Load jobs that still need review.

    Jobs are selected based on:
        - Application Status = Not Applied
        - Recommendation = APPLY or REVIEW
        - Review Decision is empty

    Returns:
        List of job dictionaries.
    """

    if not TRACKER_FILE.exists():
        print()
        print("Job tracker not found.")
        print(f"Expected file: {TRACKER_FILE}")
        print()

        return []

    workbook = load_workbook(
        TRACKER_FILE,
        read_only=True
    )

    sheet = workbook["Jobs"]

    jobs = []

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        # --------------------------------------------------
        # Skip empty rows
        # --------------------------------------------------

        if not row[1]:
            continue

        # --------------------------------------------------
        # Read columns
        # --------------------------------------------------

        recommendation = (
            row[6]
            or ""
        )

        status = (
            row[14]
            or "Not Applied"
        )

        review_decision = (
            row[15]
            or ""
        )

        # --------------------------------------------------
        # Only show jobs that need review
        # --------------------------------------------------

        if (
            status == "Not Applied"
            and recommendation in (
                "APPLY",
                "REVIEW"
            )
            and not review_decision
        ):

            jobs.append({
                "row_number": row_number,
                "date_found": row[0],
                "title": row[1],
                "company": row[2],
                "location": row[3],
                "score": row[4] or 0,
                "category": row[5],
                "recommendation": recommendation,
                "role_match": row[7],
                "matching_skills": row[8],
                "missing_skills": row[9],
                "warnings": row[10],
                "posted": row[11],
                "deadline": row[12],
                "url": row[13],
                "status": status,
                "review_decision": review_decision,
            })

    workbook.close()

    # ------------------------------------------------------
    # Highest score first
    # ------------------------------------------------------

    jobs.sort(
        key=lambda job: job["score"],
        reverse=True
    )

    return jobs


# ==========================================================
# DISPLAY JOB LIST
# ==========================================================

def display_job_list(jobs):
    """Display jobs that require review."""

    print()
    print("=" * 70)
    print("JOBS TO REVIEW")
    print("=" * 70)
    print()

    if not jobs:
        print("No jobs currently require review.")
        print()

        return

    print(
        f"Found {len(jobs)} jobs to review."
    )

    print()

    for index, job in enumerate(
        jobs,
        start=1
    ):

        print(
            f"{index}. "
            f"{job['score']}% — "
            f"{job['title']}"
        )

        print(
            f"   {job['company']} | "
            f"{job['location']} | "
            f"{job['recommendation']}"
        )

    print()


# ==========================================================
# DISPLAY JOB DETAILS
# ==========================================================

def display_job(job):
    """Display detailed information about a selected job."""

    print()
    print("=" * 70)
    print(job["title"])
    print("=" * 70)

    print()

    print(
        f"Company: {job['company']}"
    )

    print(
        f"Location: {job['location']}"
    )

    print(
        f"Score: {job['score']}%"
    )

    print(
        f"Category: {job['category']}"
    )

    print(
        f"Recommendation: {job['recommendation']}"
    )

    # ------------------------------------------------------
    # Role match
    # ------------------------------------------------------

    if job["role_match"]:

        print()

        print("Role match:")

        print(
            f"  {job['role_match']}"
        )

    # ------------------------------------------------------
    # Matching skills
    # ------------------------------------------------------

    if job["matching_skills"]:

        print()

        print("Matching skills:")

        print(
            f"  {job['matching_skills']}"
        )

    # ------------------------------------------------------
    # Missing skills
    # ------------------------------------------------------

    if job["missing_skills"]:

        print()

        print("Missing skills:")

        print(
            f"  {job['missing_skills']}"
        )

    # ------------------------------------------------------
    # Warnings
    # ------------------------------------------------------

    if job["warnings"]:

        print()

        print("Warnings:")

        print(
            f"  {job['warnings']}"
        )

    # ------------------------------------------------------
    # Posted date
    # ------------------------------------------------------

    if job["posted"]:

        print()

        print(
            f"Posted: {job['posted']}"
        )

    # ------------------------------------------------------
    # Deadline
    # ------------------------------------------------------

    if job["deadline"]:

        print(
            f"Deadline: {job['deadline']}"
        )

    # ------------------------------------------------------
    # URL
    # ------------------------------------------------------

    print()

    print("URL:")

    print(
        job["url"]
    )

    print()


# ==========================================================
# UPDATE REVIEW DECISION
# ==========================================================

def update_review_decision(
    row_number,
    decision
):
    """
    Update the Review Decision column.

    Column P = Review Decision.

    Valid decisions:
        Apply
        Maybe
        Skip
    """

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    sheet.cell(
        row=row_number,
        column=16
    ).value = decision

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()


# ==========================================================
# UPDATE APPLICATION STATUS
# ==========================================================

def update_application_status(
    row_number,
    status
):
    """
    Update the Application Status column.

    Column O = Application Status.
    """

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    sheet.cell(
        row=row_number,
        column=15
    ).value = status

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()


# ==========================================================
# JOB ACTIONS
# ==========================================================

def job_actions(job):
    """
    Show review actions for a selected job.
    """

    while True:

        print("-" * 70)

        print("What do you want to do?")

        print()

        print("1. Apply")
        print("2. Maybe")
        print("3. Skip")
        print("4. Mark as Applied")
        print("5. Back")

        print()

        choice = input(
            "Choose an option: "
        ).strip()

        # --------------------------------------------------
        # APPLY
        # --------------------------------------------------

        if choice == "1":

            update_review_decision(
                job["row_number"],
                "Apply"
            )

            update_application_status(
                job["row_number"],
                "To Apply"
            )

            print()

            print(
                "✓ Job marked as Apply."
            )

            print(
                "✓ Application status set to To Apply."
            )

            return

        # --------------------------------------------------
        # MAYBE
        # --------------------------------------------------

        elif choice == "2":

            update_review_decision(
                job["row_number"],
                "Maybe"
            )

            print()

            print(
                "✓ Job marked as Maybe."
            )

            return

        # --------------------------------------------------
        # SKIP
        # --------------------------------------------------

        elif choice == "3":

            update_review_decision(
                job["row_number"],
                "Skip"
            )

            print()

            print(
                "✓ Job marked as Skip."
            )

            return

        # --------------------------------------------------
        # MARK AS APPLIED
        # --------------------------------------------------

        elif choice == "4":

            update_review_decision(
                job["row_number"],
                "Apply"
            )

            update_application_status(
                job["row_number"],
                "Applied"
            )

            print()

            print(
                "✓ Review decision set to Apply."
            )

            print(
                "✓ Application status set to Applied."
            )

            return

        # --------------------------------------------------
        # BACK
        # --------------------------------------------------

        elif choice == "5":

            return

        # --------------------------------------------------
        # INVALID
        # --------------------------------------------------

        else:

            print()

            print(
                "Invalid option."
            )

            print(
                "Please choose 1-5."
            )

            print()


# ==========================================================
# REVIEW JOBS
# ==========================================================

def review_jobs():
    """Run the interactive job-review process."""

    while True:

        jobs = load_jobs()

        display_job_list(
            jobs
        )

        # --------------------------------------------------
        # Nothing left to review
        # --------------------------------------------------

        if not jobs:

            return

        print(
            "Enter a job number to view details."
        )

        print(
            "Enter 'q' to quit."
        )

        print()

        choice = input(
            "Select job: "
        ).strip().lower()

        # --------------------------------------------------
        # Quit
        # --------------------------------------------------

        if choice == "q":

            print()

            print(
                "Exiting job review."
            )

            return

        # --------------------------------------------------
        # Validate number
        # --------------------------------------------------

        try:

            job_number = int(
                choice
            )

        except ValueError:

            print()

            print(
                "Please enter a valid job number."
            )

            continue

        # --------------------------------------------------
        # Validate range
        # --------------------------------------------------

        if not (
            1 <= job_number <= len(jobs)
        ):

            print()

            print(
                "Invalid job number."
            )

            continue

        # --------------------------------------------------
        # Select job
        # --------------------------------------------------

        selected_job = jobs[
            job_number - 1
        ]

        # --------------------------------------------------
        # Display details
        # --------------------------------------------------

        display_job(
            selected_job
        )

        # --------------------------------------------------
        # Actions
        # --------------------------------------------------

        job_actions(
            selected_job
        )


# ==========================================================
# MAIN
# ==========================================================

if __name__ == "__main__":

    review_jobs()

