
"""
Cover letter generator.

Generates a tailored cover letter for each job using:
- Candidate profile
- Job title
- Company
- Matching skills
- Relevant experience
- Relevant projects

The generated letter is saved as a Word document.
"""

import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# CONFIGURATION
# ============================================================

PROFILE_FILE = Path("data/candidate_profile.json")
APPLICATIONS_DIR = Path("applications")


# ============================================================
# PROFILE
# ============================================================

def load_candidate_profile():
    """Load the candidate profile from JSON."""

    if not PROFILE_FILE.exists():
        raise FileNotFoundError(
            f"Candidate profile not found: {PROFILE_FILE}"
        )

    with open(
        PROFILE_FILE,
        "r",
        encoding="utf-8"
    ) as file:

        return json.load(file)


# ============================================================
# TEXT HELPERS
# ============================================================

def normalize_text(text):
    """Normalize text for keyword matching."""

    if not text:
        return ""

    text = str(text).lower()

    text = re.sub(
        r"[^a-z0-9+#.\s]",
        " ",
        text
    )

    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


def clean_sentence(text):
    """Clean a sentence extracted from profile data."""

    if not text:
        return ""

    text = text.strip()

    # Remove accidental leading punctuation
    text = re.sub(
        r"^[\s\-•]+",
        "",
        text
    )

    # Normalize whitespace
    text = re.sub(
        r"\s+",
        " ",
        text
    )

    # Fix common lowercase sentence starts
    if text:
        text = text[0].upper() + text[1:]

    # Ensure punctuation
    if text[-1] not in ".!?":
        text += "."

    return text


def get_job_text(job):
    """Return normalized job text."""

    return normalize_text(
        f"""
        {job.get('title', '')}
        {job.get('description', '')}
        {job.get('requirements', '')}
        {job.get('skills', '')}
        """
    )


# ============================================================
# ROLE DETECTION
# ============================================================

def detect_role_type(job):
    """
    Detect the broad role category.

    Returns:
        backend
        frontend
        full_stack
        software
        other
    """

    title = normalize_text(
        job.get("title", "")
    )

    if (
        "full stack" in title
        or "fullstack" in title
    ):
        return "full_stack"

    if (
        "backend" in title
        or "back end" in title
        or "python developer" in title
        or "flask developer" in title
    ):
        return "backend"

    if (
        "frontend" in title
        or "front end" in title
        or "react developer" in title
        or "web developer" in title
    ):
        return "frontend"

    if (
        "software developer" in title
        or "software engineer" in title
        or "developer" in title
        or "engineer" in title
    ):
        return "software"

    return "other"


# ============================================================
# SKILL SELECTION
# ============================================================

def select_relevant_skills(
    profile,
    job,
    match=None
):
    """
    Select candidate skills relevant to the job.

    If matcher results are available, matching skills are
    prioritized.
    """

    job_text = get_job_text(job)

    selected = {}

    technical_skills = profile.get(
        "technical_skills",
        {}
    )

    # --------------------------------------------------------
    # First use matcher results when available
    # --------------------------------------------------------

    matching_skills = []

    if match:
        matching_skills = match.get(
            "matching_skills",
            []
        )

    normalized_matches = {
        normalize_text(skill)
        for skill in matching_skills
    }

    # --------------------------------------------------------
    # Check candidate profile skills
    # --------------------------------------------------------

    for category, skills in technical_skills.items():

        relevant = []

        for skill in skills:

            normalized_skill = normalize_text(
                skill
            )

            if (
                normalized_skill in normalized_matches
                or normalized_skill in job_text
            ):
                relevant.append(skill)

        if relevant:
            selected[category] = relevant

    return selected


def flatten_skills(skills):
    """Flatten categorized skills into one list."""

    result = []

    for skill_list in skills.values():

        for skill in skill_list:

            if skill not in result:
                result.append(skill)

    return result


# ============================================================
# EXPERIENCE SELECTION
# ============================================================

def experience_relevance_score(
    experience,
    job,
    role_type
):
    """Calculate relevance of an experience item."""

    job_text = get_job_text(job)

    experience_text = normalize_text(
        f"""
        {experience.get('title', '')}
        {experience.get('company', '')}
        {' '.join(experience.get('responsibilities', []))}
        """
    )

    score = 0

    # --------------------------------------------------------
    # Direct keyword overlap
    # --------------------------------------------------------

    job_words = set(job_text.split())
    experience_words = set(
        experience_text.split()
    )

    overlap = job_words.intersection(
        experience_words
    )

    score += min(
        len(overlap) * 2,
        10
    )

    # --------------------------------------------------------
    # Software experience
    # --------------------------------------------------------

    if (
        "software" in experience_text
        or "developer" in experience_text
        or "backend" in experience_text
        or "api" in experience_text
        or "programming" in experience_text
    ):
        score += 10

    # --------------------------------------------------------
    # Backend relevance
    # --------------------------------------------------------

    if role_type == "backend":

        if (
            "python" in experience_text
            or "flask" in experience_text
            or "api" in experience_text
            or "backend" in experience_text
            or "database" in experience_text
        ):
            score += 10

    # --------------------------------------------------------
    # Frontend relevance
    # --------------------------------------------------------

    if role_type == "frontend":

        if (
            "react" in experience_text
            or "javascript" in experience_text
            or "frontend" in experience_text
            or "web" in experience_text
        ):
            score += 10

    # --------------------------------------------------------
    # Full stack relevance
    # --------------------------------------------------------

    if role_type == "full_stack":

        backend = (
            "python" in experience_text
            or "flask" in experience_text
            or "backend" in experience_text
            or "api" in experience_text
        )

        frontend = (
            "react" in experience_text
            or "javascript" in experience_text
            or "frontend" in experience_text
        )

        if backend:
            score += 7

        if frontend:
            score += 7

    return score


def select_relevant_experience(
    profile,
    job
):
    """Select the most relevant work experience."""

    role_type = detect_role_type(job)

    experience = profile.get(
        "experience",
        []
    )

    scored = []

    for item in experience:

        score = experience_relevance_score(
            item,
            job,
            role_type
        )

        scored.append(
            (
                score,
                item
            )
        )

    scored.sort(
        key=lambda item: item[0],
        reverse=True
    )

    # --------------------------------------------------------
    # Only include meaningful experience.
    #
    # For technical roles we prefer the strongest evidence
    # rather than filling the CV/letter with unrelated work.
    # --------------------------------------------------------

    selected = [
        item
        for score, item in scored
        if score >= 8
    ]

    # Maximum of two experiences in the letter
    return selected[:2]


# ============================================================
# PROJECT SELECTION
# ============================================================

def project_relevance_score(
    project,
    job,
    role_type
):
    """Calculate relevance of a project."""

    job_text = get_job_text(job)

    project_text = normalize_text(
        f"""
        {project.get('name', '')}
        {project.get('description', '')}
        {' '.join(project.get('responsibilities', []))}
        {' '.join(project.get('technologies', []))}
        """
    )

    job_words = set(
        job_text.split()
    )

    project_words = set(
        project_text.split()
    )

    overlap = job_words.intersection(
        project_words
    )

    score = min(
        len(overlap) * 3,
        15
    )

    # --------------------------------------------------------
    # Backend
    # --------------------------------------------------------

    if role_type == "backend":

        if (
            "flask" in project_text
            or "python" in project_text
            or "backend" in project_text
            or "rest api" in project_text
            or "api" in project_text
        ):
            score += 15

        if "database" in project_text:
            score += 5

    # --------------------------------------------------------
    # Frontend
    # --------------------------------------------------------

    elif role_type == "frontend":

        if (
            "react" in project_text
            or "javascript" in project_text
            or "frontend" in project_text
        ):
            score += 15

        if (
            "dashboard" in project_text
            or "web" in project_text
        ):
            score += 5

    # --------------------------------------------------------
    # Full stack
    # --------------------------------------------------------

    elif role_type == "full_stack":

        backend = (
            "flask" in project_text
            or "python" in project_text
            or "backend" in project_text
            or "api" in project_text
        )

        frontend = (
            "react" in project_text
            or "javascript" in project_text
            or "frontend" in project_text
        )

        if backend:
            score += 10

        if frontend:
            score += 10

    return score


def select_relevant_projects(
    profile,
    job
):
    """Select the strongest projects for the job."""

    role_type = detect_role_type(job)

    projects = profile.get(
        "projects",
        []
    )

    scored = []

    for project in projects:

        score = project_relevance_score(
            project,
            job,
            role_type
        )

        scored.append(
            (
                score,
                project
            )
        )

    scored.sort(
        key=lambda item: item[0],
        reverse=True
    )

    selected = [
        project
        for score, project in scored
        if score > 0
    ]

    # Keep the letter focused
    return selected[:3]


# ============================================================
# SUMMARY / OPENING
# ============================================================

def get_role_focus(role_type):
    """Return the main focus for a role."""

    if role_type == "backend":
        return (
            "backend development, REST API development, "
            "and database implementation"
        )

    if role_type == "frontend":
        return (
            "frontend development, web application development, "
            "and building user-facing interfaces"
        )

    if role_type == "full_stack":
        return (
            "full-stack web development across frontend and "
            "backend applications"
        )

    return (
        "software development and building practical "
        "web applications"
    )


def build_opening(
    job_title,
    company,
    role_type
):
    """Build the opening paragraph."""

    focus = get_role_focus(
        role_type
    )

    return (
        f"I am writing to express my interest in the "
        f"{job_title} position at {company}. As a Junior "
        f"Software Developer with hands-on experience in "
        f"{focus}, I am excited about the opportunity to "
        f"contribute my technical skills to your team."
    )


# ============================================================
# SKILL PARAGRAPH
# ============================================================

def build_skill_paragraph(
    skills,
    role_type
):
    """Build a role-specific technical skills paragraph."""

    flat_skills = flatten_skills(
        skills
    )

    if not flat_skills:
        return (
            "My technical background includes software "
            "development, web application development, and "
            "working collaboratively on technical projects."
        )

    skill_text = ", ".join(
        flat_skills
    )

    if role_type == "backend":

        return (
            f"My technical background includes {skill_text}. "
            f"I have practical experience developing backend "
            f"functionality, building REST APIs, working with "
            f"databases, and contributing to team-based "
            f"software projects."
        )

    if role_type == "frontend":

        return (
            f"My technical background includes {skill_text}. "
            f"I have practical experience building web "
            f"interfaces and applications using modern "
            f"frontend technologies and contributing to "
            f"team-based software projects."
        )

    if role_type == "full_stack":

        return (
            f"My technical background includes {skill_text}. "
            f"I have experience working across frontend and "
            f"backend development, including web interfaces, "
            f"backend services, databases, and REST APIs."
        )

    return (
        f"My technical background includes {skill_text}. "
        f"I have practical experience applying these "
        f"technologies to web applications and software "
        f"development projects."
    )


# ============================================================
# EXPERIENCE PARAGRAPH
# ============================================================

def get_best_software_experience(
    experience
):
    """Find the strongest software-related experience."""

    for item in experience:

        title = normalize_text(
            item.get("title", "")
        )

        responsibilities = normalize_text(
            " ".join(
                item.get(
                    "responsibilities",
                    []
                )
            )
        )

        combined = (
            f"{title} {responsibilities}"
        )

        if (
            "software" in combined
            or "developer" in combined
            or "backend" in combined
            or "api" in combined
        ):
            return item

    return None


def build_experience_paragraph(
    experience,
    role_type
):
    """Build a paragraph from the strongest experience."""

    if not experience:
        return (
            "My practical experience has given me an "
            "opportunity to work on software projects, "
            "collaborate with teams, and develop solutions "
            "to technical requirements."
        )

    item = get_best_software_experience(
        experience
    )

    if not item:
        item = experience[0]

    title = item.get(
        "title",
        "my previous role"
    )

    company = item.get(
        "company",
        ""
    )

    responsibilities = item.get(
        "responsibilities",
        []
    )

    cleaned = [
        clean_sentence(
            responsibility
        )
        for responsibility in responsibilities
        if responsibility
    ]

    # --------------------------------------------------------
    # Software Engineering Intern
    # --------------------------------------------------------

    if (
        "software engineering intern"
        in normalize_text(title)
    ):

        api_work = None
        collaboration = None

        for responsibility in cleaned:

            lower = normalize_text(
                responsibility
            )

            if (
                "rest api" in lower
                or "api" in lower
            ):
                api_work = responsibility

            if "collaborat" in lower:
                collaboration = responsibility

        sentences = []

        if api_work:
            sentences.append(
                f"During my experience as {title} at "
                f"{company}, I {api_work[0].lower()}"
                f"{api_work[1:]}"
            )

        if collaboration:
            sentences.append(
                f"I also {collaboration[0].lower()}"
                f"{collaboration[1:]}"
            )

        if sentences:

            if role_type == "frontend":

                return (
                    f"During my experience as {title} at "
                    f"{company}, I contributed to a production "
                    f"project and collaborated with the "
                    f"development team to deliver backend "
                    f"functionality. This strengthened my "
                    f"understanding of development workflows, "
                    f"team collaboration, and delivering "
                    f"technical work to requirements."
                )

            return (
                " ".join(sentences)
                + " This experience strengthened my ability "
                "to understand technical requirements, "
                "collaborate with development teams, and "
                "deliver reliable software functionality."
            )

    # --------------------------------------------------------
    # Generic experience fallback
    # --------------------------------------------------------

    useful = cleaned[:2]

    if useful:

        return (
            f"In my role as {title} at {company}, "
            + " ".join(useful)
            + " This experience strengthened my "
            "problem-solving, communication, and ability "
            "to contribute effectively to a team."
        )

    return (
        f"My experience as {title} at {company} has "
        "strengthened my ability to work collaboratively, "
        "understand requirements, and contribute to "
        "practical solutions."
    )


# ============================================================
# PROJECT PARAGRAPH
# ============================================================

def build_project_paragraph(
    projects,
    role_type
):
    """Build a concise project paragraph."""

    if not projects:
        return (
            "Through my software projects, I have gained "
            "practical experience applying development "
            "concepts to real-world problems."
        )

    project_names = [
        project.get(
            "name",
            ""
        )
        for project in projects
        if project.get("name")
    ]

    names = ", ".join(
        project_names
    )

    if role_type == "backend":

        return (
            f"Through projects including {names}, I have "
            f"applied my knowledge of backend development, "
            f"REST APIs, database design, and full-stack "
            f"application development. These projects have "
            f"given me practical experience working through "
            f"technical requirements and contributing to "
            f"complete software solutions."
        )

    if role_type == "frontend":

        return (
            f"Through projects including {names}, I have "
            f"applied my knowledge of frontend development, "
            f"React, JavaScript, and web application "
            f"development. These projects have strengthened "
            f"my ability to build user-facing applications "
            f"while working within broader software projects."
        )

    if role_type == "full_stack":

        return (
            f"Through projects including {names}, I have "
            f"applied my knowledge across frontend and "
            f"backend development, including React, Python, "
            f"Flask, REST APIs, and database design. These "
            f"projects have given me practical experience "
            f"contributing to complete web applications."
        )

    return (
        f"Through projects including {names}, I have "
        f"applied my software development knowledge to "
        f"practical applications, working through technical "
        f"requirements and contributing to complete "
        f"software solutions."
    )


# ============================================================
# CLOSING
# ============================================================

def build_closing():
    """Build the closing paragraph."""

    return (
        "I would welcome the opportunity to discuss how my "
        "background, technical skills, and enthusiasm for "
        "software development could contribute to your team. "
        "Thank you for considering my application. I look "
        "forward to the possibility of discussing the role "
        "with you."
    )


# ============================================================
# COVER LETTER GENERATION
# ============================================================

def generate_cover_letter(
    job,
    match=None
):
    """
    Generate a complete tailored cover letter.

    Args:
        job: Dictionary containing job information.
        match: Optional matcher result.

    Returns:
        Cover letter as a string.
    """

    profile = load_candidate_profile()

    job_title = (
        job.get(
            "title",
            "Software Developer"
        )
        .strip()
    )

    company = (
        job.get(
            "company",
            "your organization"
        )
        .strip()
    )

    role_type = detect_role_type(
        job
    )

    # --------------------------------------------------------
    # Select relevant information
    # --------------------------------------------------------

    skills = select_relevant_skills(
        profile,
        job,
        match
    )

    experience = select_relevant_experience(
        profile,
        job
    )

    projects = select_relevant_projects(
        profile,
        job
    )

    # --------------------------------------------------------
    # Build paragraphs
    # --------------------------------------------------------

    opening = build_opening(
        job_title,
        company,
        role_type
    )

    skill_paragraph = build_skill_paragraph(
        skills,
        role_type
    )

    experience_paragraph = build_experience_paragraph(
        experience,
        role_type
    )

    project_paragraph = build_project_paragraph(
        projects,
        role_type
    )

    closing = build_closing()

    # --------------------------------------------------------
    # Assemble letter
    # --------------------------------------------------------

    return (
        "Dear Hiring Manager,\n\n"
        f"{opening}\n\n"
        f"{skill_paragraph}\n\n"
        f"{experience_paragraph}\n\n"
        f"{project_paragraph}\n\n"
        f"{closing}\n\n"
        "Kind regards,\n"
        f"{profile['personal']['name']}"
    )


# ============================================================
# WORD DOCUMENT
# ============================================================

def create_document():
    """Create and configure a Word document."""

    document = Document()

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    normal.paragraph_format.space_after = Pt(6)

    return document


def add_letter_header(
    document,
    profile
):
    """Add candidate contact information."""

    personal = profile["personal"]

    # --------------------------------------------------------
    # Name
    # --------------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        personal["name"]
    )

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)

    # --------------------------------------------------------
    # Contact
    # --------------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    contact = (
        f"{personal['phone']} | "
        f"{personal['email']} | "
        f"{personal['location']}"
    )

    run = paragraph.add_run(
        contact
    )

    run.font.name = "Arial"
    run.font.size = Pt(9)

    # --------------------------------------------------------
    # Links
    # --------------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    links = (
        f"{personal['linkedin']} | "
        f"{personal['github']}"
    )

    run = paragraph.add_run(
        links
    )

    run.font.name = "Arial"
    run.font.size = Pt(9)


def add_letter_body(
    document,
    letter
):
    """Add cover letter text to the document."""

    paragraphs = letter.split(
        "\n\n"
    )

    for text in paragraphs:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(8)

        run = paragraph.add_run(
            text
        )

        run.font.name = "Arial"
        run.font.size = Pt(10.5)


# ============================================================
# FILE NAME
# ============================================================

def create_application_directory_name(
    job
):
    """
    Create a safe application directory name.
    """

    title = job.get(
        "title",
        "software-developer"
    )

    company = job.get(
        "company",
        "company"
    )

    name = (
        f"{company}-{title}"
    )

    name = normalize_text(
        name
    )

    name = re.sub(
        r"[^a-z0-9]+",
        "-",
        name
    )

    name = name.strip(
        "-"
    )

    return name[:120]


# ============================================================
# SAVE COVER LETTER
# ============================================================

def save_cover_letter(
    job,
    match=None,
    output_directory=None
):
    """
    Generate and save a tailored cover letter.

    Returns:
        Path to generated DOCX.
    """

    profile = load_candidate_profile()

    letter = generate_cover_letter(
        job,
        match
    )

    document = create_document()

    add_letter_header(
        document,
        profile
    )

    # Space between header and letter
    document.add_paragraph()

    add_letter_body(
        document,
        letter
    )

    # --------------------------------------------------------
    # Determine output directory
    # --------------------------------------------------------

    if output_directory is None:

        output_directory = (
            APPLICATIONS_DIR
            / create_application_directory_name(
                job
            )
        )

    output_directory = Path(
        output_directory
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True
    )

    output_file = (
        output_directory
        / "cover_letter.docx"
    )

    document.save(
        output_file
    )

    return output_file


# ============================================================
# DISPLAY HELPERS
# ============================================================

def print_cover_letter(
    job,
    letter,
    skills,
    experience,
    projects
):
    """Display generated cover letter information."""

    print()
    print("=" * 60)
    print(
        f"Job: {job.get('title', '')}"
    )
    print(
        f"Company: {job.get('company', '')}"
    )
    print()

    print("Matching skills:")

    flat_skills = flatten_skills(
        skills
    )

    if flat_skills:

        print(
            "  "
            + ", ".join(
                flat_skills
            )
        )

    else:

        print("  None")

    print()

    print("Selected experience:")

    if experience:

        for item in experience:

            print(
                f"  {item.get('title', '')} "
                f"at "
                f"{item.get('company', '')}"
            )

    else:

        print("  None")

    print()

    print("Selected projects:")

    if projects:

        for project in projects:

            print(
                f"  {project.get('name', '')}"
            )

    else:

        print("  None")

    print()

    print("Generated cover letter:")
    print("-" * 60)
    print(letter)
    print("-" * 60)


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    print("=" * 60)
    print("COVER LETTER TEST")
    print("=" * 60)

    profile = load_candidate_profile()

    test_jobs = [

        {
            "title": "Backend Developer",
            "company": "Test Company",
            "description": """
            We are looking for a backend developer with
            Python, Flask, REST APIs, SQL and Git experience.
            """
        },

        {
            "title": "Frontend Developer",
            "company": "Test Company",
            "description": """
            We are looking for a frontend developer with
            HTML, CSS, JavaScript, React.js and Next.js
            experience.
            """
        },

        {
            "title": "Full Stack Developer",
            "company": "Test Company",
            "description": """
            We are looking for a full stack developer with
            JavaScript, React.js, Python, Flask, REST APIs,
            SQL and Git experience.
            """
        }
    ]

    for job in test_jobs:

        skills = select_relevant_skills(
            profile,
            job
        )

        experience = select_relevant_experience(
            profile,
            job
        )

        projects = select_relevant_projects(
            profile,
            job
        )

        letter = generate_cover_letter(
            job
        )

        print_cover_letter(
            job,
            letter,
            skills,
            experience,
            projects
        )

    # --------------------------------------------------------
    # Generate one test DOCX
    # --------------------------------------------------------

    output = save_cover_letter(
        test_jobs[0],
        output_directory=(
            APPLICATIONS_DIR
            / "test-application"
        )
    )

    print()
    print(
        f"Created cover letter: {output}"
    )

    print("=" * 60)
    print("COVER LETTER TEST COMPLETE")
    print("=" * 60)

