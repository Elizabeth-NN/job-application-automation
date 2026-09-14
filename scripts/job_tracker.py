
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


TRACKER_FILE = Path("data/job_tracker.xlsx")


# ==========================================================
# TRACKER SCHEMA
# ==========================================================

HEADERS = [
    "Date Found",
    "Job Title",
    "Company",
    "Location",
    "Score",
    "Category",
    "Recommendation",
    "Role Match",
    "Matching Skills",
    "Missing Skills",
    "Warnings",
    "Posted",
    "Deadline",
    "URL",

    # Application information
    "Application Method",
    "Application Email",
    "Application Subject",
    "Application URL",

    # Application tracking
    "Application Status",
    "Review Decision",
    "CV Version",
    "Cover Letter",
    "Notes",
]


# ==========================================================
# COLUMN NUMBERS
# ==========================================================

COLUMN = {
    "date_found": 1,
    "job_title": 2,
    "company": 3,
    "location": 4,
    "score": 5,
    "category": 6,
    "recommendation": 7,
    "role_match": 8,
    "matching_skills": 9,
    "missing_skills": 10,
    "warnings": 11,
    "posted": 12,
    "deadline": 13,
    "url": 14,

    "application_method": 15,
    "application_email": 16,
    "application_subject": 17,
    "application_url": 18,

    "application_status": 19,
    "review_decision": 20,
    "cv_version": 21,
    "cover_letter": 22,
    "notes": 23,
}


# ==========================================================
# COLORS
# ==========================================================

STRONG_FILL = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

GOOD_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9"
)

POSSIBLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC"
)

WEAK_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6"
)

POOR_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)


# ==========================================================
# HELPERS
# ==========================================================

def _as_list(value):
    """Convert a value into a list of strings."""

    if value is None:
        return []

    if isinstance(value, str):
        return [value] if value else []

    return list(value)


def _join_values(values, separator=", "):
    """Safely convert values into a string."""

    return separator.join(
        str(value)
        for value in _as_list(values)
        if value
    )


# ==========================================================
# TRACKER FORMATTING
# ==========================================================

def _apply_header_formatting(sheet):
    """Apply formatting to the header row."""

    for cell in sheet[1]:
        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


def _apply_data_validation(sheet):
    """Add dropdowns for application tracking fields."""

    # Remove existing validations before recreating them.
    sheet.data_validations.dataValidation = []

    # ------------------------------------------------------
    # APPLICATION STATUS
    # ------------------------------------------------------

    status_dropdown = DataValidation(
        type="list",
        formula1=(
            '"Not Applied,To Apply,Applied,'
            'Interview,Rejected,Offer,Withdrawn"'
        ),
        allow_blank=False
    )

    sheet.add_data_validation(
        status_dropdown
    )

    status_dropdown.add(
        f"{get_column_letter(COLUMN['application_status'])}2:"
        f"{get_column_letter(COLUMN['application_status'])}1000"
    )

    # ------------------------------------------------------
    # REVIEW DECISION
    # ------------------------------------------------------

    review_dropdown = DataValidation(
        type="list",
        formula1='"Apply,Maybe,Skip"',
        allow_blank=True
    )

    sheet.add_data_validation(
        review_dropdown
    )

    review_dropdown.add(
        f"{get_column_letter(COLUMN['review_decision'])}2:"
        f"{get_column_letter(COLUMN['review_decision'])}1000"
    )


def _apply_column_widths(sheet):
    """Set readable widths for tracker columns."""

    widths = {
        1: 14,
        2: 45,
        3: 30,
        4: 20,
        5: 10,
        6: 20,
        7: 18,
        8: 30,
        9: 45,
        10: 45,
        11: 50,
        12: 15,
        13: 15,
        14: 70,

        # Application information
        15: 22,
        16: 35,
        17: 45,
        18: 70,

        # Application tracking
        19: 20,
        20: 18,
        21: 35,
        22: 35,
        23: 40,
    }

    for column, width in widths.items():
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width


def _apply_default_alignment(sheet):
    """Apply default alignment to all populated cells."""

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


# ==========================================================
# TRACKER MIGRATION
# ==========================================================

def _migrate_tracker_schema():
    """
    Update an existing tracker to the current schema.

    Existing data is preserved by matching columns by header name.
    New columns are added automatically.
    """

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    existing_headers = [
        cell.value
        for cell in sheet[1]
    ]

    # ------------------------------------------------------
    # Check whether migration is necessary
    # ------------------------------------------------------

    if existing_headers == HEADERS:
        _apply_header_formatting(sheet)
        _apply_column_widths(sheet)
        _apply_default_alignment(sheet)
        _apply_data_validation(sheet)

        sheet.freeze_panes = "A2"

        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(HEADERS))}1"
        )

        workbook.save(TRACKER_FILE)
        workbook.close()

        return

    # ------------------------------------------------------
    # Preserve existing data by header name
    # ------------------------------------------------------

    old_data = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        old_data.append(
            dict(
                zip(
                    existing_headers,
                    row
                )
            )
        )

    # ------------------------------------------------------
    # Replace sheet with current schema
    # ------------------------------------------------------

    workbook.remove(sheet)

    sheet = workbook.create_sheet(
        "Jobs",
        0
    )

    sheet.append(HEADERS)

    # ------------------------------------------------------
    # Restore existing rows
    # ------------------------------------------------------

    for old_row in old_data:

        new_row = [
            old_row.get(
                header,
                ""
            )
            for header in HEADERS
        ]

        sheet.append(
            new_row
        )

    # ------------------------------------------------------
    # Format migrated tracker
    # ------------------------------------------------------

    _apply_header_formatting(sheet)

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(HEADERS))}1"
    )

    _apply_data_validation(sheet)

    _apply_column_widths(sheet)

    _apply_default_alignment(sheet)

    # ------------------------------------------------------
    # Format existing rows
    # ------------------------------------------------------

    for row_number in range(
        2,
        sheet.max_row + 1
    ):
        format_job_row(
            sheet,
            row_number
        )

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    print(
        "Tracker schema updated successfully."
    )


# ==========================================================
# TRACKER CREATION
# ==========================================================

def create_tracker():
    """Create or migrate the Excel tracker."""

    TRACKER_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # ------------------------------------------------------
    # Create new tracker
    # ------------------------------------------------------

    if not TRACKER_FILE.exists():

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Jobs"

        sheet.append(
            HEADERS
        )

        _apply_header_formatting(
            sheet
        )

        sheet.freeze_panes = "A2"

        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(HEADERS))}1"
        )

        _apply_data_validation(
            sheet
        )

        _apply_column_widths(
            sheet
        )

        _apply_default_alignment(
            sheet
        )

        workbook.save(
            TRACKER_FILE
        )

        workbook.close()

        print(
            f"Created tracker: {TRACKER_FILE}"
        )

        return

    # ------------------------------------------------------
    # Existing tracker
    # ------------------------------------------------------

    _migrate_tracker_schema()


# ==========================================================
# FIND JOB
# ==========================================================

def find_job_row(url):
    """
    Find a job row using its URL.

    Returns:
        Row number if found.
        None if not found.
    """

    if not url:
        return None

    create_tracker()

    workbook = load_workbook(
        TRACKER_FILE,
        read_only=True
    )

    sheet = workbook["Jobs"]

    try:

        for row_number in range(
            2,
            sheet.max_row + 1
        ):

            existing_url = sheet.cell(
                row=row_number,
                column=COLUMN["url"]
            ).value

            if existing_url == url:
                return row_number

    finally:

        workbook.close()

    return None


# ==========================================================
# JOB EXISTS
# ==========================================================

def job_exists(url):
    """Check whether a job already exists."""

    if not url:
        return False

    return (
        find_job_row(url)
        is not None
    )


# ==========================================================
# FIND ROW NUMBER
# ==========================================================

def get_job_row(url):
    """
    Return the actual Excel row number for a job URL.

    Returns:
        Row number if found.
        None if not found.
    """

    return find_job_row(url)


# ==========================================================
# CATEGORY FORMATTING
# ==========================================================

def apply_category_formatting(
    sheet,
    row_number
):
    """Apply color formatting based on match category."""

    category = sheet.cell(
        row=row_number,
        column=COLUMN["category"]
    ).value

    if not category:
        return

    category = str(
        category
    ).upper()

    fill = None

    if category == "STRONG MATCH":
        fill = STRONG_FILL

    elif category == "GOOD MATCH":
        fill = GOOD_FILL

    elif category == "POSSIBLE MATCH":
        fill = POSSIBLE_FILL

    elif category == "WEAK MATCH":
        fill = WEAK_FILL

    elif category == "POOR MATCH":
        fill = POOR_FILL

    if fill:

        for column in range(
            1,
            len(HEADERS) + 1
        ):

            sheet.cell(
                row=row_number,
                column=column
            ).fill = fill


# ==========================================================
# FORMAT ROW
# ==========================================================

def format_job_row(
    sheet,
    row_number
):
    """Apply formatting to a job row."""

    for column in range(
        1,
        len(HEADERS) + 1
    ):

        cell = sheet.cell(
            row=row_number,
            column=column
        )

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

    # ------------------------------------------------------
    # Category colors
    # ------------------------------------------------------

    apply_category_formatting(
        sheet,
        row_number
    )

    # ------------------------------------------------------
    # Score alignment
    # ------------------------------------------------------

    score_cell = sheet.cell(
        row=row_number,
        column=COLUMN["score"]
    )

    score_cell.alignment = Alignment(
        horizontal="center",
        vertical="top"
    )

    # ------------------------------------------------------
    # Job URL hyperlink
    # ------------------------------------------------------

    url_cell = sheet.cell(
        row=row_number,
        column=COLUMN["url"]
    )

    url = url_cell.value

    if url:

        url_cell.hyperlink = url
        url_cell.style = "Hyperlink"

    # ------------------------------------------------------
    # Application URL hyperlink
    # ------------------------------------------------------

    application_url_cell = sheet.cell(
        row=row_number,
        column=COLUMN["application_url"]
    )

    application_url = (
        application_url_cell.value
    )

    if application_url:

        application_url_cell.hyperlink = (
            application_url
        )

        application_url_cell.style = (
            "Hyperlink"
        )


# ==========================================================
# SAVE NEW JOB
# ==========================================================

def save_job(
    job,
    match_result
):
    """
    Save a matched job to the Excel tracker.

    If the job already exists, it is not duplicated.

    Returns:
        True  -> new job was added
        False -> job already existed
    """

    create_tracker()

    url = job.get(
        "url",
        ""
    )

    # ======================================================
    # DUPLICATE CHECK
    # ======================================================

    if job_exists(url):

        print(
            "Already in tracker."
        )

        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    # ======================================================
    # MATCH DATA
    # ======================================================

    role_matches = match_result.get(
        "role_matches",
        []
    )

    matching_skills = match_result.get(
        "matching_skills",
        []
    )

    missing_skills = match_result.get(
        "missing_skills",
        []
    )

    warnings = match_result.get(
        "warnings",
        []
    )

    # ======================================================
    # APPLICATION DATA
    # ======================================================

    application_method = job.get(
        "application_method",
        job.get(
            "method",
            ""
        )
    )

    application_email = job.get(
        "application_email",
        job.get(
            "email",
            ""
        )
    )

    application_subject = job.get(
        "application_subject",
        job.get(
            "subject",
            ""
        )
    )

    application_url = job.get(
        "application_url",
        ""
    )

    # ======================================================
    # ADD JOB
    # ======================================================

    sheet.append([

        # --------------------------------------------------
        # Job information
        # --------------------------------------------------

        datetime.now().strftime(
            "%Y-%m-%d"
        ),

        job.get(
            "title",
            ""
        ),

        job.get(
            "company",
            ""
        ),

        job.get(
            "location",
            ""
        ),

        # --------------------------------------------------
        # Matching information
        # --------------------------------------------------

        match_result.get(
            "score",
            0
        ),

        match_result.get(
            "category",
            ""
        ),

        match_result.get(
            "recommendation",
            ""
        ),

        _join_values(
            role_matches
        ),

        _join_values(
            matching_skills
        ),

        _join_values(
            missing_skills
        ),

        _join_values(
            warnings,
            separator=" | "
        ),

        # --------------------------------------------------
        # Job metadata
        # --------------------------------------------------

        job.get(
            "posted",
            ""
        ),

        job.get(
            "deadline",
            ""
        ),

        url,

        # --------------------------------------------------
        # Application information
        # --------------------------------------------------

        application_method,

        application_email,

        application_subject,

        application_url,

        # --------------------------------------------------
        # Application tracking
        # --------------------------------------------------

        "Not Applied",

        "",

        "",

        "",

        "",
    ])

    row_number = sheet.max_row

    format_job_row(
        sheet,
        row_number
    )

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True


# ==========================================================
# UPDATE APPLICATION INFORMATION
# ==========================================================

def update_application_info(
    url,
    method=None,
    email=None,
    subject=None,
    application_url=None
):
    """
    Update application information for an existing job.

    This is useful when application information is extracted
    separately from the initial job collection process.
    """

    if not url:
        return False

    create_tracker()

    row_number = get_job_row(
        url
    )

    if row_number is None:
        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    # ------------------------------------------------------
    # Method
    # ------------------------------------------------------

    if method is not None:

        sheet.cell(
            row=row_number,
            column=COLUMN["application_method"]
        ).value = method

    # ------------------------------------------------------
    # Email
    # ------------------------------------------------------

    if email is not None:

        sheet.cell(
            row=row_number,
            column=COLUMN["application_email"]
        ).value = email

    # ------------------------------------------------------
    # Subject
    # ------------------------------------------------------

    if subject is not None:

        sheet.cell(
            row=row_number,
            column=COLUMN["application_subject"]
        ).value = subject

    # ------------------------------------------------------
    # Application URL
    # ------------------------------------------------------

    if application_url is not None:

        cell = sheet.cell(
            row=row_number,
            column=COLUMN["application_url"]
        )

        cell.value = application_url

        if application_url:

            cell.hyperlink = application_url
            cell.style = "Hyperlink"

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True


# ==========================================================
# UPDATE APPLICATION DOCUMENTS
# ==========================================================

def update_application_documents(
    url,
    cv_path=None,
    cover_letter_path=None
):
    """
    Update CV and cover-letter paths for an existing job.

    Document generation happens after the job has already
    been saved to the tracker.
    """

    if not url:
        return False

    create_tracker()

    row_number = get_job_row(
        url
    )

    if row_number is None:
        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    # ======================================================
    # CV PATH
    # ======================================================

    if cv_path:

        sheet.cell(
            row=row_number,
            column=COLUMN["cv_version"]
        ).value = str(
            Path(cv_path)
        )

    # ======================================================
    # COVER LETTER PATH
    # ======================================================

    if cover_letter_path:

        sheet.cell(
            row=row_number,
            column=COLUMN["cover_letter"]
        ).value = str(
            Path(cover_letter_path)
        )

    # ======================================================
    # APPLICATION STATUS
    # ======================================================

    if cv_path or cover_letter_path:

        sheet.cell(
            row=row_number,
            column=COLUMN["application_status"]
        ).value = "To Apply"

    # ======================================================
    # FORMAT DOCUMENT CELLS
    # ======================================================

    for column in (
        COLUMN["cv_version"],
        COLUMN["cover_letter"]
    ):

        cell = sheet.cell(
            row=row_number,
            column=column
        )

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True


# ==========================================================
# UPDATE NOTES
# ==========================================================

def update_job_notes(
    url,
    notes
):
    """Update the Notes column for an existing job."""

    if not url:
        return False

    create_tracker()

    row_number = get_job_row(
        url
    )

    if row_number is None:
        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    sheet.cell(
        row=row_number,
        column=COLUMN["notes"]
    ).value = notes or ""

    sheet.cell(
        row=row_number,
        column=COLUMN["notes"]
    ).alignment = Alignment(
        vertical="top",
        wrap_text=True
    )

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True


# ==========================================================
# UPDATE APPLICATION STATUS
# ==========================================================

def update_application_status(
    url,
    status
):
    """Update the application status for an existing job."""

    if not url:
        return False

    create_tracker()

    row_number = get_job_row(
        url
    )

    if row_number is None:
        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    sheet.cell(
        row=row_number,
        column=COLUMN["application_status"]
    ).value = status

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True


# ==========================================================
# UPDATE REVIEW DECISION
# ==========================================================

def update_review_decision(
    url,
    decision
):
    """Update the manual review decision for a job."""

    if not url:
        return False

    create_tracker()

    row_number = get_job_row(
        url
    )

    if row_number is None:
        return False

    workbook = load_workbook(
        TRACKER_FILE
    )

    sheet = workbook["Jobs"]

    sheet.cell(
        row=row_number,
        column=COLUMN["review_decision"]
    ).value = decision or ""

    workbook.save(
        TRACKER_FILE
    )

    workbook.close()

    return True
